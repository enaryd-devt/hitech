# -*- encoding: utf-8 -*-

from odoo import api, fields, models, _
from odoo.exceptions import ValidationError
import base64
import io
import logging

from openpyxl import Workbook
from openpyxl.styles import Font, Border, Side, Alignment
from openpyxl.utils import get_column_letter

logger = logging.getLogger(__name__)


class HrPayrollDipe(models.TransientModel):
    _name = 'hr.payroll.dipe'

    period_id = fields.Many2one('hr.pai.period', string='Période de paie')
    state = fields.Selection([('choose', 'Choisir'), ('get', 'Obtenir')], default='choose')
    fname = fields.Char("Nom du fichier", readonly=True)
    data = fields.Binary('Fichier', readonly=True)

    def get_payroll_dipe(self):
        buffer = io.BytesIO()

        wb = Workbook()
        ws = wb.active
        ws.title = _('Dipe')

        # ---------------- STYLES ----------------
        thin = Side(style='thin')
        border = Border(left=thin, right=thin, top=thin, bottom=thin)

        style = {
            "font": Font(name="Calibri", size=11),
            "border": border
        }
        style_bold = {
            "font": Font(name="Calibri", size=11, bold=True),
            "border": border
        }
        style_bold10 = {
            "font": Font(name="Calibri", size=10, bold=True),
            "border": border,
            "alignment": Alignment(horizontal="center", vertical="center")
        }
        style_bold12 = {
            "font": Font(name="Calibri", size=12, bold=True),
            "border": border,
            "alignment": Alignment(horizontal="center", vertical="center")
        }
        style_bold14 = {
            "font": Font(name="Calibri", size=14, bold=True),
            "border": border,
            "alignment": Alignment(horizontal="center", vertical="center")
        }

        def apply_style(cell, st):
            cell.font = st.get("font")
            cell.border = st.get("border")
            if st.get("alignment"):
                cell.alignment = st.get("alignment")

        # ---------------- COLUMN WIDTHS ----------------
        ws.column_dimensions['A'].width = 29
        payslip_list = self.env['hr.payslip'].search([
            ('period_id', '=', self.period_id.id),
            ('state', '!=', 'cancel')
        ]).sorted(lambda p: p.employee_id.matricule)

        for i in range(2, len(payslip_list) + 3):
            ws.column_dimensions[get_column_letter(i)].width = 20

        # ---------------- TITLE ----------------
        ws.merge_cells(start_row=1, start_column=2,
                       end_row=1, end_column=2 + len(payslip_list))
        cell = ws.cell(row=1, column=2, value="DIPE orick")
        apply_style(cell, style_bold14)

        # ---------------- HEADERS ----------------
        header1 = [
            "Noms & Prénoms", "Date d'embauche", "Matricule interne", "Matricule CNPS", "Cat/Ech",
            "Salaire de base", "Sursalaire", "Prime d'ancienneté", "Prime de responsabilité",
            "Prime de suggestion", "Prime de caisse", "Prime particulière",
            "Ind. De représentation", "Ind. De logement", "Ind. D'eléctricité",
            "Ind. D'eau", "Ind. De transport", "Ind. De panier", "Congés payés", "",
            "Salaire brut", "Salaire taxable intermediare", "Salaire taxable",
            "Salaire cotisable", "", "Cotisations fiscales", "", "IRPP",
            "CAC/IRPP", "CFC", "RAV", "TDL", "CFC", "FNE", "", "Totaux", "",
            "Cotisations sociales", "", "AT/MP", "AF", "PVID", "", "Totaux", "",
            "Salaire net"
        ]

        start_row = 3
        for i, head in enumerate(header1):
            c = ws.cell(row=start_row + i, column=1, value=head)
            apply_style(c, style_bold)

        # ---------------- RULE GROUPS ----------------
        indem1 = ["BASE", "SURSAL", "PAN", "INDRESP", "INDSUGG", "PRITENUCAI",
                  "PRIEXECPT", "INDREP", "INDLOG", "ELECTRICITE", "EAU",
                  "INDEPLT", "INPANIER", "INDEMCONGES"]

        indem2 = ["BRUT", "SALTAX", "SALTAX_B1", "SALCO_PLAF"]
        indem3 = ["IRPPH", "CAC", "CFC", "RAV", "TC", "CFC_PAT", "FNE"]
        indem4 = ["AT_MP", "PRESTFAM", "PV_PAT"]

        def get_amount_by_rule_code(rules, slip):
            return [
                sum(l.total for l in slip.line_ids.filtered(lambda x: x.code == r))
                for r in rules
            ]

        # ---------------- DATA ----------------
        col = 2
        for slip in payslip_list:
            row = start_row
            cat = f"{slip.category_id.name or ''}/{slip.echelon_id.name or ''}"

            values = [
                slip.employee_id.name,
                slip.employee_id.first_contract_date or "",
                slip.employee_id.matricule or "",
                slip.employee_id.cnps or "",
                cat
            ]
            values += get_amount_by_rule_code(indem1, slip)
            values += [""]
            values += get_amount_by_rule_code(indem2, slip)
            values += ["", "", ""]
            values += get_amount_by_rule_code(indem3, slip)
            values += ["", sum(get_amount_by_rule_code(indem3, slip)), "", "", ""]
            values += get_amount_by_rule_code(indem4, slip)
            values += ["", sum(get_amount_by_rule_code(indem4, slip)), ""]
            values += get_amount_by_rule_code(["NET"], slip)

            for val in values:
                cell = ws.cell(row=row, column=col, value=val)
                apply_style(cell, style)
                row += 1

            col += 1

        # ---------------- TOTAL COLUMN ----------------
        def total_rule(rule):
            return sum(
                sum(l.total for l in s.line_ids.filtered(lambda x: x.code == rule))
                for s in payslip_list
            )

        totals = ["Totaux", "", "", "", ""]
        totals += [total_rule(r) for r in indem1]
        totals += [""]
        totals += [total_rule(r) for r in indem2]
        totals += ["", "", ""]
        t_fisc = sum(total_rule(r) for r in indem3)
        totals += [total_rule(r) for r in indem3]
        totals += ["", t_fisc, "", "", ""]
        t_soc = sum(total_rule(r) for r in indem4)
        totals += [total_rule(r) for r in indem4]
        totals += ["", t_soc, ""]
        totals += [total_rule("NET")]

        row = start_row
        for val in totals:
            cell = ws.cell(row=row, column=col, value=val)
            apply_style(cell, style_bold)
            row += 1

        # ---------------- SAVE ----------------
        wb.save(buffer)
        buffer.seek(0)

        self.data = base64.b64encode(buffer.read())
        self.fname = _('Dipe ORICK %s.xlsx') % self.period_id.code
        self.state = 'get'

        view = self.env.ref('primetech_rh.hr_payroll_dipe_view_form')
        return {
            'name': _('Exporter la DIPE'),
            'type': 'ir.actions.act_window',
            'view_mode': 'form',
            'res_model': self._name,
            'res_id': self.id,
            'view_id': view.id,
            'target': 'new'
        }
